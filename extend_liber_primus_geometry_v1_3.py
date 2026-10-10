from pathlib import Path
import openpyxl, json, math, hashlib
import numpy as np
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE=Path('/mnt/data')
inp=BASE/'liber_primus_visual_master_dataset_v1_2.xlsx'
out=BASE/'liber_primus_visual_master_dataset_v1_3.xlsx'
report_in=BASE/'LIBER_PRIMUS_VISUAL_GEOMETRY_RESEARCH_v1_2.md'
report_out=BASE/'LIBER_PRIMUS_VISUAL_GEOMETRY_RESEARCH_v1_3.md'
wb=openpyxl.load_workbook(inp)
if 'Mirror Audit v1.3' in wb.sheetnames: del wb['Mirror Audit v1.3']
ws=wb.create_sheet('Mirror Audit v1.3')
rows=[
['Plant ornament mirror audit — v1.3'],
['Reanalysis of stored Auxiliary Controls v1.2 records only. No raw JPEG pixels were independently decoded in this pass.'],
[],
['Finding ID','Question','Status','Result','Interpretation / limit'],
['MIR-01','Page 57 vs page 68 right-side plant group','Exact stored-record match','All five corresponding labels have identical centroid x/y, component area, and bounding boxes. Max absolute coordinate delta = 0 px.','This is equality of workbook records, not an independent image-pixel comparison. It may reflect reuse of a common ornament or copied measurements.'],
['MIR-02','Left-side partial candidates vs reflected right-side points','Three-point near-reflection','For 3 correspondences, x_left = (2400 - x_right) - 1 px and y_left = y_right exactly at stored precision.','This supports bilateral layout symmetry in the recorded geometry; the left side is incomplete (3 candidates vs 5 points), so it is not a complete five-point reflection test.'],
['MIR-03','Cipher interpretation','Not supported','No coordinate-to-rune mapping or text prediction is tested.','Treat as decorative/template reuse control until a preregistered, falsifiable cipher relationship is specified.'],
[],
['Page 57 right label','x right px','y right px','Reflected x = 2400 − x','Observed left candidate','Observed left x px','Observed left y px','dx = left_x − reflected_x','dy = left_y − right_y'],
]
# stored data
src=wb['Auxiliary Controls v1.2']
records=list(src.iter_rows(min_row=1,values_only=True))
right={}
left=[]
for r in records:
    if not isinstance(r[0],(int,float)): continue
    page=int(r[0]); group=r[1]; label=str(r[2])
    if group not in ('plant-side right five-dot group','plant-side left ROI — incomplete'): continue
    x=float(r[5]); y=float(r[6]); area=r[7]; box=(r[8],r[9],r[10],r[11])
    if page==57 and group=='plant-side right five-dot group': right[label]=(x,y,area,box)
    if page==57 and group=='plant-side left ROI — incomplete': left.append((label,x,y))
# closest matches by y; only exact y match for known rows
for lab in ['label_3','label_2','label_4']:
    x,y,area,box=right[lab]
    target=2400-x
    match=min(left,key=lambda q:abs(q[2]-y)+abs(q[1]-target))
    _,lx,ly=match
    rows.append([lab,x,y,target,match[0],lx,ly,lx-target,ly-y])
for r in rows: ws.append(r)
ws.freeze_panes='A5'; ws.auto_filter.ref=f'A4:E7'
for cell in ws[1]: cell.font=Font(bold=True,size=14)
for cell in ws[4]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='24364B'); cell.alignment=Alignment(wrap_text=True,vertical='top')
for row in ws.iter_rows():
    for c in row:
        c.alignment=Alignment(vertical='top',wrap_text=True)
for i in range(1,10): ws.column_dimensions[get_column_letter(i)].width=[18,28,28,42,45,20,20,26,24][i-1]
ws.row_dimensions[1].height=24; ws.row_dimensions[2].height=35
wb.save(out)

# Compute an explicit machine-readable summary from the recorded data
p57={str(k):{'x':v[0],'y':v[1],'area':v[2],'bbox':list(v[3])} for k,v in right.items()}
p68={}
for r in records:
    if isinstance(r[0],(int,float)) and int(r[0])==68 and r[1]=='plant-side right five-dot group':
        p68[str(r[2])]={'x':float(r[5]),'y':float(r[6]),'area':r[7],'bbox':[r[8],r[9],r[10],r[11]]}
summary={'version':'1.3','input_workbook':inp.name,'output_workbook':out.name,
 'scope':'derived only from stored Auxiliary Controls v1.2; no source-pixel decode',
 'page57_page68_record_equality': all(p57[k]==p68.get(k) for k in p57) and set(p57)==set(p68),
 'max_abs_centroid_delta_px': max(max(abs(p57[k]['x']-p68[k]['x']),abs(p57[k]['y']-p68[k]['y'])) for k in p57),
 'left_reflection_matches':[],
 'interpretation':'record-level evidence consistent with repeated ornament / bilateral symmetry; not cipher evidence'}
for lab in ['label_3','label_2','label_4']:
    x,y,area,box=right[lab]; rx=2400-x
    match=min(left,key=lambda q:abs(q[2]-y)+abs(q[1]-rx))
    summary['left_reflection_matches'].append({'right_label':lab,'right_xy':[x,y],'reflected_x':rx,'left_candidate':match[0],'left_xy':[match[1],match[2]],'dx_px':match[1]-rx,'dy_px':match[2]-y})
json_out=BASE/'liber_primus_visual_geometry_audit_v1_3.json'
json_out.write_text(json.dumps(summary,indent=2))

old=report_in.read_text()
addition='''\n\n## v1.3 addendum — plant ornament symmetry control\n\n**Purpose:** Continue the existing audit with a narrow, falsification-oriented check of the plant-side ornament records on pages 57 and 68. This is a derivation from the stored `Auxiliary Controls v1.2` rows, not a new JPEG decode.\n\nThe five right-side plant-group points on page 57 and the five corresponding points on page 68 have exactly equal recorded centroids, areas, and bounding boxes. This makes the repeated-ornament/template explanation the leading baseline to test before assigning cryptographic meaning. The coordinate equality alone cannot tell whether the source artwork is pixel-identical, whether measurements were copied between rows, or whether the pages independently reproduce the same design.\n\nA second check compared the three stored left-side candidates to the horizontal reflection of the page-57 right group around the page centerline, using `x_reflected = 2400 - x`. The three closest correspondences are:\n\n| Right-side point | Reflected x | Left candidate x | dx (left − reflected) | y difference |\n|---|---:|---:|---:|---:|\n| label_3 | 367.155 | 366.155 | −1.000 px | 0.000 px |\n| label_2 | 424.979 | 423.979 | −1.000 px | 0.000 px |\n| label_4 | 420.438 | 419.438 | −1.000 px | 0.000 px |\n\nThis is a three-point, incomplete left-side check, not a full five-point correspondence. The exact shared y values and consistent one-pixel x offset are compatible with a mirrored ornament or mirrored layout, but the evidence remains record-level. Do not treat the one-pixel offset as an encoded value without an independently justified rule.\n\n**Updated conclusion:** this control strengthens the need to model repeated and mirrored decoration in the manuscript-empirical null. It does not strengthen the case for a cipher relationship. The next valid measurement step is still to recover source JPEG bytes and decoder provenance, then perform an independent pixel comparison of the corresponding page-57/page-68 crops and rerun the detector across all 75 page images.\n'''
report_out.write_text(old+addition)
print('WROTE',out, out.stat().st_size)
print('WROTE',report_out, report_out.stat().st_size)
print('WROTE',json_out, json_out.stat().st_size)
print(json.dumps(summary,indent=2))
