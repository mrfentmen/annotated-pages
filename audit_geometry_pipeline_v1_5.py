from pathlib import Path
import openpyxl, hashlib, json, subprocess, sys
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

base=Path('/mnt/data')
source=base/'files/measure_constellations_v0_2.py'
master=base/'liber_primus_visual_master_dataset_v1_4.xlsx'
out=base/'liber_primus_visual_master_dataset_v1_5.xlsx'
report=base/'LIBER_PRIMUS_VISUAL_GEOMETRY_RESEARCH_v1_5.md'
audit_json=base/'liber_primus_geometry_pipeline_audit_v1_5.json'
code=source.read_text()
lines=code.splitlines()
def line_no(fragment):
    for i,l in enumerate(lines,1):
        if fragment in l: return i
    return None

# Verify the existing synthetic self-test (does not access source scans).
p=subprocess.run([sys.executable,str(source),'--self-test'],capture_output=True,text=True,check=False)
try: selftest=json.loads(p.stdout)
except Exception: selftest={'stdout':p.stdout,'stderr':p.stderr,'returncode':p.returncode}

findings=[
['PIPE-01','Scope','Confirmed limitation','Only three predefined target ROIs (pages 24, 40, 73) are processed.','PAGE_CONFIG and run() loop; this is a target-group extractor, not a corpus-wide detector.','Do not use this script to claim all manuscript dot groups were searched.'],
['PIPE-02','Source validation','Good guard, untested locally','Git blob SHA-1 and 2400×3600 dimensions are checked before measurement.','git_blob_sha1() and run() source checks.','Retain these guards; execute only after exact bytes are available.'],
['PIPE-03','Rejected components','Output omission','Components that fail the dot-like filter are discarded, not emitted to an audit table; comments imply broader candidate reporting.','extract_components() filter and return path.','Modify extractor to log all components with accepted/rejected status and rejection reasons.'],
['PIPE-04','Geometric signature','Method mismatch','pairwise_signature() divides distances by the median; workbook Cluster Signatures v1.2 says distances are normalized by RMS pairwise distance.','pairwise_signature() near top of script vs workbook methodology.','Do not compare script outputs numerically with workbook signatures unless normalization is unified.'],
['PIPE-05','Null model','Invalid for inference','Random points in a rectangle ignore real manuscript marks, layouts, copied ornament, component detection, and target-selection procedure. Script can emit empirical p-values despite its own caveat.','run() random control section.','Disable inferential interpretation; use only as software smoke-test or historical exploratory diagnostic.'],
['PIPE-06','Threshold matching','Greedy matching','Cross-threshold centroids are paired greedily by nearest distance, not via globally optimal assignment; in crowded ROIs this can create mismatches.','threshold stability block in run().','Use a minimum-cost bipartite assignment and preserve unmatched candidates.'],
['PIPE-07','Detector bias','Pre-filter selection','Area, width, height, aspect, and fill cutoffs are tuned to solid circular marks; irregular, touching, faded, or small marks are excluded before any corpus-level comparison.','extract_components() candidate filter.','For corpus survey, retain every component and classify later; calibrate on labeled controls.'],
['PIPE-08','Self-test','Pass, limited scope','Synthetic self-test passed for thresholds 40, 60, 80, 100 with five synthetic circles and zero rounded-center error.','Executed during v1.5 audit.','Validates a narrow component-extraction behavior only; not source images, JPEG decoding, ROIs, or cipher relevance.'],
['PIPE-09','Pixel decoder','Independent cross-check pending','This script uses Pillow Image.convert("L"), whereas recorded centroid sheet attributes values to a separate JavaScript baseline-JPEG luma decoder.','Pillow conversion in run(); Decoder Crosscheck v1.2.','Re-run both decoders on identical verified bytes; compare luma arrays and component centroids.'],
['PIPE-10','Pairwise function','Input assumptions unchecked','rms_difference() uses zip and would silently truncate mismatched signatures; callers currently expect ten distances from five points.','rms_difference().','Add equal-length assertions and report candidate count before any signature calculation.'],
['PIPE-11','Output reproducibility','Partial','CSV fields are generated from accepted candidates only; rejected-component evidence and full pixel-decoder settings are absent.','candidate CSV construction.','Store full detector config, software versions, all components, and accepted/rejected reasons.'],
['PIPE-12','Current research implication','No new cipher evidence','Static review and synthetic self-test do not independently validate the five-dot match, page 57/68 ornament identity, or any cipher relationship.','Scope of this audit.','Next evidence gate remains source-byte acquisition followed by independent decode and blind full-corpus detection.'],
]
summary={
'version':'1.5',
'script':source.name,
'script_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),
'source_pixels_processed':False,
'synthetic_self_test_returncode':p.returncode,
'synthetic_self_test':selftest,
'findings':[
 {'id':r[0],'area':r[1],'status':r[2],'finding':r[3],'evidence':r[4],'next_step':r[5]} for r in findings
],
'conclusion':'Static pipeline audit found several reproducibility/inference gaps. The synthetic self-test passed, but no original Liber Primus source pixels were processed and no geometric or cipher result was independently reproduced.'
}
audit_json.write_text(json.dumps(summary,indent=2))

wb=openpyxl.load_workbook(master)
if 'Pipeline Audit v1.5' in wb.sheetnames: del wb['Pipeline Audit v1.5']
ws=wb.create_sheet('Pipeline Audit v1.5')
ws.append(['Measurement pipeline audit — v1.5'])
ws.append(['Static review of measure_constellations_v0_2.py plus execution of its synthetic-only self-test. No original source JPEG bytes were processed in this pass.'])
ws.append([])
ws.append(['Finding ID','Area','Status','Finding','Code / evidence','Required next step'])
for r in findings: ws.append(r)
ws.freeze_panes='A5'; ws.auto_filter.ref=f'A4:F{4+len(findings)}'
for c in ws[1]: c.font=Font(bold=True,size=14)
for c in ws[4]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='24364B'); c.alignment=Alignment(wrap_text=True,vertical='top')
widths=[14,24,27,66,55,60]
for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
for row in ws.iter_rows():
    for c in row: c.alignment=Alignment(vertical='top',wrap_text=True)
ws.row_dimensions[1].height=24; ws.row_dimensions[2].height=36
for i in range(5,5+len(findings)): ws.row_dimensions[i].height=78
wb.save(out)

text=f'''# Liber Primus Visual Geometry Investigation — Research Record v1.5

**Date:** 2026-10-10  
**Master dataset:** `liber_primus_visual_master_dataset_v1_5.xlsx` (v1.4 retained, one audit sheet appended)  
**Pipeline reviewed:** `measure_constellations_v0_2.py`  
**Source-pixel measurements made in this pass:** No

## Result of this round

I continued with the next available reproducibility task: a static audit of the existing centroid-extraction pipeline and an execution of its built-in synthetic self-test. The self-test passed at grayscale thresholds 40, 60, 80, and 100, finding all five synthetic circles with zero rounded-center error. This validates only that narrow synthetic case. It does **not** validate the original JPEG decoder, selected regions of interest, actual dot centroids, repeated artwork, or a cipher connection.

The exact source image bytes still are not present in the local runtime. An attempt to retrieve the raw GitHub image from the runtime failed because external DNS/network access is unavailable there. The public Git blob inventory in v1.4 establishes repository object identity, not independent pixel reproduction.

## Main pipeline audit findings

1. **The script is not a full-corpus detector.** It only processes preselected regions on pages 24, 40, and 73. It cannot support claims about the absence or frequency of similar geometry elsewhere in the 75-page manuscript.
2. **Rejected components are discarded.** The filter keeps only components that look like solid circular dots by area, bounding box, aspect ratio, and fill ratio. Rejected components and rejection reasons are not exported, which prevents auditing detector bias.
3. **The signature normalization differs from the workbook.** The script divides pair distances by the median distance; `Cluster Signatures v1.2` documents RMS pair-distance normalization. These numbers are not directly comparable.
4. **The optional random null is not suitable for inferential claims.** It samples points uniformly in a rectangle rather than modeling the actual page artwork, mark distribution, repeated/mirrored ornaments, detection process, and target-selection process. The script's computed empirical p-values must not be presented as evidence of significance.
5. **Threshold matching is greedy.** It pairs components across thresholds by nearest available centroid. With crowded or irregular regions, greedy pairing can fail to find the globally best one-to-one assignment and may misstate centroid stability.
6. **The decoder cross-check remains open.** The script uses Pillow's grayscale conversion. The stored centroid sheet says its measurements came from a separate JavaScript baseline-JPEG luma decoder. Those outputs need to be compared on identical verified JPEG bytes before either is treated as independently validated.
7. **Signature length is not checked.** The RMS-difference helper uses `zip`, so unequal-length vectors would be silently truncated instead of rejected. Current five-point targets are expected to produce ten distances, but the assumption should be enforced.

## Current interpretation

This audit does not change the geometric evidence. The previously recorded near-identical five-dot signatures for pages 24, 40, and 73 remain a promising descriptive observation, but their original-pixel measurements are not independently reproduced here. The stored page 57/68 equality likewise remains record-level evidence, not a fresh pixel comparison. Neither result establishes a cipher relationship.

## Next evidence gate

1. Acquire the exact original JPEG bytes for pages 24, 40, 57, 68, and 73 (then all pages) in an environment that can write those bytes to disk.
2. Verify each file against its Git blob SHA-1 before decoding.
3. Run Pillow/libjpeg and the recorded JavaScript luma decoder on the exact same bytes; compare decoded luma pixels or, at minimum, independently extracted centroids and connected-component masks.
4. Preserve **all** connected components and their rejection reasons; freeze detector parameters before running the 75-page scan.
5. Blindly review candidate groups, identify copied-artwork families, and only then construct a manuscript-empirical control distribution that repeats the complete target-selection procedure.
6. Keep cipher relevance as a separate test requiring a predeclared coordinate/order-to-rune prediction. Geometry similarity alone is not enough.

## Files

- `liber_primus_visual_master_dataset_v1_5.xlsx` — authoritative master with `Pipeline Audit v1.5` appended.
- `liber_primus_geometry_pipeline_audit_v1_5.json` — machine-readable findings and synthetic self-test result.
- `audit_geometry_pipeline_v1_5.py` — script used to generate this audit.

The v1.4 source integrity inventory and earlier research sheets are retained unchanged in the master workbook.
'''
report.write_text(text)
print(json.dumps({'master':str(out),'master_bytes':out.stat().st_size,'report':str(report),'report_bytes':report.stat().st_size,'audit_json':str(audit_json),'audit_json_bytes':audit_json.stat().st_size,'script_sha256':summary['script_sha256'],'selftest_returncode':p.returncode,'selftest':selftest},indent=2))
